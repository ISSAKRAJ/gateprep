import openpyxl
import json
from datetime import datetime

wb = openpyxl.load_workbook('GATE_CS_2027_Study_Tracker.xlsx', data_only=True)

data = {
    'summary': {},
    'subject_progress': [],
    'weekly_summary': [],
    'tracker': [],
    'mocks': [],
    'mistakes': [],
    'syllabus': []
}

# 1. Summary sheet
s_sum = wb['Summary']
data['summary'] = {
    'title': s_sum.cell(1, 1).value,
    'instructions': s_sum.cell(2, 1).value,
    'exam_date': '2027-02-06',
    'total_days': 104,
    'total_planned_hours': 660,
    'notes': s_sum.cell(52, 1).value
}

# Subjects progress (R22 to R32)
for r in range(22, 33):
    subj = s_sum.cell(r, 1).value
    if subj:
        data['subject_progress'].append({
            'subject': subj,
            'days_planned': s_sum.cell(r, 2).value or 0,
            'days_done': s_sum.cell(r, 3).value or 0,
            'percent_done': s_sum.cell(r, 4).value or 0,
            'confidence': s_sum.cell(r, 5).value or 3
        })

# Weekly summary (R36 to R50)
for r in range(36, 51):
    w_date = s_sum.cell(r, 1).value
    if w_date:
        w_str = w_date.strftime('%Y-%m-%d') if hasattr(w_date, 'strftime') else str(w_date)
        data['weekly_summary'].append({
            'week_starting': w_str,
            'planned_hrs': s_sum.cell(r, 2).value or 0,
            'actual_hrs': s_sum.cell(r, 3).value or 0,
            'days_done': s_sum.cell(r, 4).value or 0,
            'pyqs_solved': s_sum.cell(r, 5).value or 0
        })

# 2. Tracker sheet
s_track = wb['Tracker']
headers = [s_track.cell(1, c).value for c in range(1, s_track.max_column+1)]
for r in range(2, s_track.max_row+1):
    vals = [s_track.cell(r, c).value for c in range(1, s_track.max_column+1)]
    if any(v is not None for v in vals) and vals[0] is not None:
        d = {}
        for h, v in zip(headers, vals):
            if hasattr(v, 'strftime'):
                v = v.strftime('%Y-%m-%d')
            d[h] = v
        data['tracker'].append(d)

# 3. Mocks
s_mocks = wb['Mocks']
mock_headers = [s_mocks.cell(2, c).value for c in range(1, s_mocks.max_column+1)]
for r in range(3, s_mocks.max_row+1):
    vals = [s_mocks.cell(r, c).value for c in range(1, s_mocks.max_column+1)]
    if any(v is not None for v in vals) and vals[0] is not None:
        d = {}
        for h, v in zip(mock_headers, vals):
            if hasattr(v, 'strftime'):
                v = v.strftime('%Y-%m-%d')
            d[h] = v
        data['mocks'].append(d)

# 4. Mistake Log
s_mistakes = wb['Mistake Log']
m_headers = [s_mistakes.cell(2, c).value for c in range(1, s_mistakes.max_column+1)]
for r in range(3, s_mistakes.max_row+1):
    vals = [s_mistakes.cell(r, c).value for c in range(1, s_mistakes.max_column+1)]
    if any(v is not None for v in vals) and vals[1] is not None:
        d = {}
        for h, v in zip(m_headers, vals):
            if hasattr(v, 'strftime'):
                v = v.strftime('%Y-%m-%d')
            d[h] = v
        data['mistakes'].append(d)

# 5. Syllabus
s_syl = wb['Syllabus 2027']
for r in range(3, s_syl.max_row+1):
    subj = s_syl.cell(r, 1).value
    changed = s_syl.cell(r, 2).value
    todo = s_syl.cell(r, 3).value
    if subj and not str(subj).startswith('Source:'):
        data['syllabus'].append({
            'subject': subj,
            'changed': changed,
            'todo': todo
        })

with open('tracker_data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, ensure_ascii=False)

print('SUCCESS')
